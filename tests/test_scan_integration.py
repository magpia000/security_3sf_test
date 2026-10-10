"""scan 통합 테스트(태스크 5.1~5.4).

fixture 응답으로 collect를 실행해 묶음을 만든 뒤, 네트워크를 막은 상태에서 scan한다.
"""

import json
import time
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook

from oss_vuln_fix_finder.__main__ import main
from oss_vuln_fix_finder.collect.runner import run_collect
from oss_vuln_fix_finder.constants import KST, OUTPUT_COLUMNS
from oss_vuln_fix_finder.mapping import load_mapping
from oss_vuln_fix_finder.scan import run_scan
from tests.collect.fakes import FakeResponse
from tests.collect.test_collectors import fixture_router, make_client
from tests.conftest import PROJECT_ROOT

MAPPING = PROJECT_ROOT / "config" / "component_map.yaml"
COLLECTED = datetime(2026, 10, 10, 15, 0, tzinfo=KST)
SCANNED = COLLECTED + timedelta(hours=1)
INPUT_ROWS = [
    (245, "zlib", "1.2.5"),
    (52, "OpenSSL", "0.9.8h"),
    (37, "googlegson", "1.7"),
    (244, "zlib", "1.1.3"),
    (900, "unknown-lib", "1.0"),
    (175, "SQLite", "3.7.9"),
    (901, "OpenSSL", None),
    (902, "OpenSSL", "3.6.0"),
]


@pytest.fixture
def scan_setup(tmp_path, make_input_xlsx):
    """입력 엑셀과 묶음(SQLite는 수집 실패)을 만든다."""
    input_path = make_input_xlsx(INPUT_ROWS)

    def handler(method, url, kwargs):
        if (kwargs.get("params") or {}).get("virtualMatchString") == "cpe:2.3:a:sqlite:sqlite":
            return FakeResponse(400)
        return fixture_router(method, url, kwargs)

    client, _ = make_client(handler)
    bundle = run_collect(
        input_path, tmp_path / "bundles", load_mapping(MAPPING), client=client, now=COLLECTED
    )
    return input_path, bundle.bundle_path, tmp_path


def _rows(path):
    sheet = load_workbook(path).active
    return [tuple(cell.value for cell in row) for row in sheet.iter_rows()]


def _scan(setup, **kwargs):
    input_path, bundle_path, tmp_path = setup
    return run_scan(
        input_path,
        bundle_path,
        tmp_path / "out",
        load_mapping(MAPPING),
        log_dir=tmp_path / "logs",
        now=SCANNED,
        **kwargs,
    )


def test_report_content(scan_setup, no_network) -> None:
    """출력 열·정렬·상태 행·조치 버전이 PRD를 따른다."""
    result = _scan(scan_setup)
    rows = _rows(result.report_path)
    assert rows[0] == OUTPUT_COLUMNS
    body = rows[1:]

    assert [r[0] for r in body] == sorted((r[0] for r in body), key=int)  # No 오름차순
    gson = [r for r in body if r[1] == "googlegson"]
    assert [r[3] for r in gson] == [
        "CVE-2022-25647",
        "GHSA-0000-0000-0000",
    ]  # CVSS 내림차순, 점수 없음은 뒤
    assert gson[0][5:7] == (7.5, "HIGH")  # OSV 전용 CVE의 점수는 NVD 단건 조회로 보완
    assert gson[0][7] == "=1.7; <2.8.9"
    assert gson[0][11] == "OSV"
    assert gson[0][8:10] == ("1.7.1", "2.11.0")

    openssl = [r for r in body if r[0] == "52"]
    assert {r[3] for r in openssl} == {"CVE-2009-0001", "CVE-2014-0224", "CVE-2016-0800"}
    assert openssl[0][8:10] == ("0.9.8zh", "3.6.0")
    assert "LTS: 3.5.4" in openssl[0][10] and "버전 목록 출처: NVD CPE" in openssl[0][10]
    assert openssl[0][12] == COLLECTED.isoformat(timespec="seconds")

    zlib_125 = [r for r in body if r[0] == "245"]
    assert {r[8] for r in zlib_125} == {"1.2.13"}
    zlib_113 = [r for r in body if r[0] == "244"]
    assert zlib_113[0][8] == "1.2.13" and "브랜치 변경" in zlib_113[0][10]

    by_no = {r[0]: r for r in body}
    assert by_no["900"][3] == "UNMAPPED"
    assert by_no["175"][3] == "UNCOLLECTED"
    assert by_no["901"][3] == "INVALID_INPUT"
    assert by_no["902"][3] == "취약점 없음" and by_no["902"][8:10] == ("-", "3.6.0")


def test_summary_and_log(scan_setup, no_network) -> None:
    """콘솔 요약 집계와 JSONL 로그."""
    result = _scan(scan_setup)
    summary = result.summary
    assert summary.input_rows == 8
    assert summary.status_counts == {
        "UNMAPPED": 1,
        "UNCOLLECTED": 1,
        "INVALID_INPUT": 1,
        "취약점 없음": 1,
    }
    assert summary.vulnerable_rows == 4
    lines = [json.loads(line) for line in result.log_path.read_text(encoding="utf-8").splitlines()]
    assert lines[-1]["type"] == "summary" and len(lines) == 9


def test_include_disputed_option(scan_setup, no_network) -> None:
    """--include-disputed를 주면 disputed CVE도 나온다."""
    rows = _rows(_scan(scan_setup, include_disputed=True).report_path)
    assert any(r[3] == "CVE-2010-5298" for r in rows)


def test_rerun_is_identical_and_fast(scan_setup, no_network) -> None:
    """같은 묶음으로 두 번 실행하면 결과가 같고(조회일시 제외), 30초 안에 끝난다."""
    started = time.monotonic()
    first = _rows(_scan(scan_setup).report_path)
    second = _rows(_scan(scan_setup).report_path)
    assert time.monotonic() - started < 30
    assert [r[:12] for r in first] == [r[:12] for r in second]


def test_warnings_for_changed_input_and_old_bundle(scan_setup, make_input_xlsx, no_network) -> None:
    """입력이 수집 때와 다르거나 묶음이 30일을 넘으면 경고한다."""
    input_path, bundle_path, tmp_path = scan_setup
    changed = make_input_xlsx([(1, "zlib", "1.2.5")])
    result = run_scan(
        changed,
        bundle_path,
        tmp_path / "out",
        load_mapping(MAPPING),
        log_dir=tmp_path / "logs",
        now=COLLECTED + timedelta(days=31),
    )
    assert len(result.summary.warnings) == 2


def test_scan_cli(scan_setup, no_network, capsys) -> None:
    """scan 명령은 보고서 경로와 요약을 출력한다."""
    input_path, bundle_path, tmp_path = scan_setup
    code = main(
        [
            "scan",
            "--input",
            str(input_path),
            "--bundle",
            str(bundle_path),
            "--output",
            str(tmp_path / "o"),
            "--mapping",
            str(MAPPING),
            "--log-dir",
            str(tmp_path / "l"),
        ]
    )
    output = capsys.readouterr().out
    assert code == 0
    assert "보고서:" in output and "취약점 발견 4행" in output


def test_scan_does_not_import_collect() -> None:
    """scan 경로 모듈은 collect(네트워크) 패키지를 import하지 않는다(rules.md 5장)."""
    import subprocess
    import sys

    code = (
        "import sys, oss_vuln_fix_finder.scan; "
        "bad = [m for m in sys.modules "
        "if m.startswith('oss_vuln_fix_finder.collect') or m == 'requests']; "
        "print(bad); sys.exit(1 if bad else 0)"
    )
    completed = subprocess.run(
        [sys.executable, "-c", code], cwd=PROJECT_ROOT, capture_output=True, text=True
    )
    assert completed.returncode == 0, completed.stdout


def test_every_report_row_has_thirteen_columns() -> None:
    """상태 행·취약점 없음 행·취약점 행 모두 13열이다."""
    from oss_vuln_fix_finder.models import (
        ComponentResult,
        InputRow,
        Remediation,
        RowStatus,
        Vulnerability,
    )
    from oss_vuln_fix_finder.scan import build_report_rows

    row = InputRow("1", "x", "1.0")
    results = [
        ComponentResult(row, RowStatus.UNMAPPED),
        ComponentResult(row, RowStatus.NO_VULNERABILITY, remediation=Remediation("-", "2.0")),
        ComponentResult(row, None, [Vulnerability("CVE-1")], remediation=Remediation("1.1", "2.0")),
    ]
    rows = build_report_rows(results, "t")
    assert {len(r) for r in rows} == {len(OUTPUT_COLUMNS)}
    assert rows[1][8:10] == ("-", "2.0") and rows[2][8:10] == ("1.1", "2.0")
