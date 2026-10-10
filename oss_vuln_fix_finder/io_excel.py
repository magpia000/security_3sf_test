"""엑셀 입출력.

입력 컴포넌트 목록을 읽고(PRD 4.1-2, 4-1), 분석 결과 보고서를 쓴다(PRD 4.5, 6장).
모든 입력 값은 문자열로 다룬다(rules.md 4.3).
"""

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from oss_vuln_fix_finder.constants import (
    INPUT_HEADER_NAME,
    INPUT_HEADER_NO,
    INPUT_HEADER_VERSION,
    INPUT_SHEET,
    OUTPUT_COLUMNS,
    OUTPUT_FILE_PREFIX,
)
from oss_vuln_fix_finder.models import InputRow

logger = logging.getLogger(__name__)


class InputFormatError(ValueError):
    """입력 엑셀에 필수 열이 없을 때 발생한다."""


def read_components(path: Path) -> list[InputRow]:
    """입력 엑셀을 읽어 InputRow 목록을 만든다.

    완전히 빈 행은 건너뛰고, 컴포넌트명이나 버전이 빈 행은 `is_valid=False`로 남긴다.

    Raises:
        InputFormatError: 필수 열(No, 컴포넌트명, 버전)을 찾지 못했을 때.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if INPUT_SHEET in workbook.sheetnames:
            sheet = workbook[INPUT_SHEET]
        else:
            sheet = workbook.worksheets[0]
            logger.warning("'%s' 시트가 없어 첫 시트 '%s'를 읽는다", INPUT_SHEET, sheet.title)
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise InputFormatError(f"{path}: 헤더 행이 없다")
        columns = _locate_columns(header)
        return [row for row in (_to_input_row(values, columns) for values in rows) if row]
    finally:
        workbook.close()


def _locate_columns(header: tuple) -> dict[str, int]:
    """헤더에서 필수 열의 위치를 찾는다. 줄바꿈 뒤 영문 병기는 무시한다."""
    wanted = {
        INPUT_HEADER_NO: INPUT_HEADER_NO,
        INPUT_HEADER_NAME: INPUT_HEADER_NAME,
        INPUT_HEADER_VERSION: INPUT_HEADER_VERSION,
    }
    found: dict[str, int] = {}
    for index, cell in enumerate(header):
        title = to_text(cell).split("\n")[0].strip()
        for key, prefix in wanted.items():
            if key not in found and title.startswith(prefix):
                found[key] = index
    missing = [key for key in wanted if key not in found]
    if missing:
        raise InputFormatError(f"입력 엑셀에 필수 열이 없다: {', '.join(missing)}")
    return found


def _to_input_row(values: tuple, columns: dict[str, int]) -> InputRow | None:
    """엑셀 한 행을 InputRow로 바꾼다. 완전히 빈 행이면 None."""

    def cell(key: str) -> str:
        index = columns[key]
        return to_text(values[index]) if index < len(values) else ""

    row_no, name, version = (
        cell(INPUT_HEADER_NO),
        cell(INPUT_HEADER_NAME),
        cell(INPUT_HEADER_VERSION),
    )
    if not (row_no or name or version):
        return None
    return InputRow(row_no=row_no, name=name, version=version, is_valid=bool(name and version))


def to_text(value: object) -> str:
    """셀 값을 문자열로 바꾼다.

    숫자 셀은 엑셀이 실수로 저장하므로 `13.0`처럼 되지 않게 정수 값은 정수로 쓴다.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# ---------- 출력 (PRD 4.5) ----------

SEVERITY_FILLS = {
    "CRITICAL": ("8B0000", "FFFFFF"),  # (배경색, 글자색) 진한 빨강
    "HIGH": ("FF0000", "FFFFFF"),
    "MEDIUM": ("FFA500", "000000"),
    "LOW": ("FFFF00", "000000"),
}
SEVERITY_COLUMN = 7  # 1부터 센 열 번호
WRAP_COLUMNS = frozenset({5, 8, 11})  # 설명, 영향 범위, 비고
MAX_COLUMN_WIDTH = 60
MIN_COLUMN_WIDTH = 8


def write_report(rows: list[tuple], output_dir: Path, now: datetime) -> Path:
    """보고서 엑셀을 `<접두어>_YYYYMMDD_HHMM.xlsx`로 저장하고 경로를 돌려준다.

    Args:
        rows: 출력 열 순서(OUTPUT_COLUMNS)에 맞춘 행 목록. 정렬은 호출하는 쪽에서 끝낸다.
        output_dir: 저장 폴더(없으면 만든다).
        now: 파일 이름에 쓸 실행 시각.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET
    sheet.append(list(OUTPUT_COLUMNS))
    for row in rows:
        sheet.append(list(row))
    _apply_format(sheet)

    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{OUTPUT_FILE_PREFIX}_{now.strftime('%Y%m%d_%H%M')}.xlsx"
    counter = 2
    while path.exists():
        path = output_dir / f"{OUTPUT_FILE_PREFIX}_{now.strftime('%Y%m%d_%H%M')}_{counter}.xlsx"
        counter += 1
    workbook.save(path)
    return path


def _apply_format(sheet) -> None:
    """헤더 굵게·고정, 열 너비, 줄바꿈, 심각도 색을 적용한다(PRD 6장)."""
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A2"

    for column_index, column in enumerate(sheet.iter_cols(min_row=1), start=1):
        longest = max(
            (len(str(cell.value)) for cell in column if cell.value is not None), default=0
        )
        width = min(max(longest + 2, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
        if column_index in WRAP_COLUMNS:
            for cell in column[1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    for cell in sheet.iter_rows(min_row=2, min_col=SEVERITY_COLUMN, max_col=SEVERITY_COLUMN):
        colors = SEVERITY_FILLS.get(str(cell[0].value or "").upper())
        if colors:
            background, font_color = colors
            cell[0].fill = PatternFill("solid", start_color=background, end_color=background)
            cell[0].font = Font(color=font_color, bold=True)
